"""Read the merged deck back and print a quick summary."""
import io
import sys
from pathlib import Path
from pptx import Presentation

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

OUT = Path(r"C:\Automation\reports\Scaling Quality at BSG_merged.pptx")
prs = Presentation(str(OUT))
print(f"Slides: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    title = ""
    if slide.shapes.title and slide.shapes.title.has_text_frame:
        title = slide.shapes.title.text_frame.text
    n_ph = len(list(slide.placeholders))
    notes = ""
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text
    notes_preview = notes.replace("\n", " | ")[:90]
    print(f"  [{i:02d}] layout={slide.slide_layout.name!r:<26} "
          f"title={title!r:<50} placeholders={n_ph} "
          f"notes={notes_preview!r}")
