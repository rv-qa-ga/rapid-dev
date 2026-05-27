import io
import sys
from pathlib import Path
from pptx import Presentation

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

OUT = Path(r"C:\Automation\reports\Accelerant_Scaling Quality at BSG.pptx")
prs = Presentation(str(OUT))
print(f"File: {OUT.name}")
print(f"Size: {OUT.stat().st_size/1024:.1f} KB")
print(f"Slides: {len(prs.slides)}")
for i, slide in enumerate(prs.slides):
    title = ""
    if slide.shapes.title and slide.shapes.title.has_text_frame:
        title = slide.shapes.title.text_frame.text
    body_lines = 0
    body_preview = ""
    for shape in slide.shapes:
        if shape.is_placeholder and shape.placeholder_format.type is not None:
            if "TITLE" in str(shape.placeholder_format.type):
                continue
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            if para.text.strip():
                body_lines += 1
                if not body_preview:
                    body_preview = para.text[:70]
    has_notes = slide.has_notes_slide and bool(slide.notes_slide.notes_text_frame.text.strip())
    print(f"  [{i:02d}] layout={slide.slide_layout.name!r:<22} title={title[:55]!r:<57} "
          f"body_lines={body_lines} preview={body_preview!r} notes={'Y' if has_notes else '-'}")
