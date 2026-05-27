"""Merge Word-doc content into a PowerPoint template.

Takes a .docx containing 'Slide N — Title' markers with bullets/notes under each,
and appends one slide per marker to a .pptx template, using the template's
'Two Content' layout (left side filled, right placeholder removed). Speaker cues
('->>' emoji, 'Say this:', 'Anchor line ...:', 'Speaker Notes', etc.) are routed
to each slide's speaker notes.
"""
from __future__ import annotations

import io
import re
import sys
from pathlib import Path

from docx import Document
from pptx import Presentation

TEMPLATE = Path(r"C:\Automation\docs\ppt\Scaling Quality at BSG.pptx")
CONTENT = Path(r"C:\Automation\docs\ppt\Offsite Presentation Content.docx")
OUTPUT = Path(r"C:\Automation\reports\Scaling Quality at BSG_merged.pptx")

BODY_LAYOUT_NAME = "Two Content"

# "Slide 1 - Title", "Slide 12 – Title" (em/en dash or hyphen)
SLIDE_MARKER = re.compile(
    r"^\s*Slide\s+(\d+)\s*[\u2014\u2013-]\s*(.+?)\s*$",
    re.IGNORECASE,
)

CUE_FINGER = "\U0001F449"  # the pointing-hand emoji used as a speaker cue
CUE_LABEL_WORDS = {"anchor line", "say this", "pause and say", "close", "speaker notes"}
SECTION_WORDS_RE = re.compile(
    r"^(?:\d+\.\s+opening|opening|closing|section\s+\d+)\b",
    re.IGNORECASE,
)


def _strip_lead(text: str) -> str:
    """Strip leading whitespace, punctuation, and symbols/emoji from the front of a line."""
    return re.sub(r"^[\s\W_]+", "", text, flags=re.UNICODE)


def is_section_marker(text: str) -> bool:
    return bool(SECTION_WORDS_RE.match(_strip_lead(text).strip()))


def is_cue_label_only(text: str) -> bool:
    core = _strip_lead(text).strip()
    core = re.sub(r"\s*\([^)]*\)\s*", " ", core)  # drop parentheticals
    core = core.rstrip(":").strip().lower()
    return core in CUE_LABEL_WORDS


def is_cue_gate(text: str) -> bool:
    if text.lstrip().startswith(CUE_FINGER):
        return True
    return is_cue_label_only(text)


def parse_doc(path: Path) -> list[dict]:
    """Walk the docx and group paragraphs into slide specs."""
    doc = Document(str(path))
    slides: list[dict] = []
    cur: dict | None = None
    in_notes = False
    next_indent = 0

    for p in doc.paragraphs:
        raw = p.text.replace("\n", " ")
        text = raw.strip()

        if not text:
            next_indent = 0
            continue

        m = SLIDE_MARKER.match(text)
        if m:
            if cur:
                slides.append(cur)
            cur = {
                "num": int(m.group(1)),
                "title": m.group(2).strip(),
                "bullets": [],
                "notes": [],
            }
            in_notes = False
            next_indent = 0
            continue

        if is_section_marker(text):
            continue

        if cur is None:
            continue

        if is_cue_gate(text):
            in_notes = True
            if is_cue_label_only(text):
                continue
            cleaned = text.lstrip(CUE_FINGER).strip().lstrip(":").strip()
            cur["notes"].append(cleaned)
            continue

        if in_notes:
            cur["notes"].append(text)
            continue

        level = next_indent
        cur["bullets"].append((level, text))
        if re.match(r"^Q[1-4]$", text) or text.endswith(":"):
            next_indent = 1
        else:
            next_indent = level

    if cur:
        slides.append(cur)
    return slides


def find_layout(prs: Presentation, name: str):
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name == name:
                return layout
    available = [ly.name for m in prs.slide_masters for ly in m.slide_layouts]
    raise KeyError(f"Layout {name!r} not found. Available: {available}")


def get_placeholder(slide, idx: int):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def remove_placeholder(slide, idx: int) -> None:
    ph = get_placeholder(slide, idx)
    if ph is None:
        return
    sp = ph._element
    sp.getparent().remove(sp)


def set_title(slide, text: str) -> None:
    title_ph = slide.shapes.title
    if title_ph is None:
        title_ph = get_placeholder(slide, 0)
    if title_ph is not None and title_ph.has_text_frame:
        title_ph.text_frame.text = text


def fill_bullets(placeholder, bullets: list[tuple[int, str]]) -> None:
    tf = placeholder.text_frame
    tf.word_wrap = True
    if not bullets:
        tf.text = ""
        return
    first_level, first_text = bullets[0]
    tf.text = first_text
    tf.paragraphs[0].level = first_level
    for level, text in bullets[1:]:
        para = tf.add_paragraph()
        para.text = text
        para.level = level


def set_notes(slide, lines: list[str]) -> None:
    if not lines:
        return
    tf = slide.notes_slide.notes_text_frame
    tf.text = "\n".join(lines)


def merge() -> Path:
    specs = parse_doc(CONTENT)
    print(f"Parsed {len(specs)} slides from Word doc:")
    for s in specs:
        print(
            f"  Slide {s['num']:>2}: {s['title']!r:<55} "
            f"bullets={len(s['bullets'])} notes={len(s['notes'])}"
        )

    prs = Presentation(str(TEMPLATE))
    existing = len(prs.slides)
    body_layout = find_layout(prs, BODY_LAYOUT_NAME)

    for s in specs:
        slide = prs.slides.add_slide(body_layout)
        set_title(slide, s["title"])
        left = get_placeholder(slide, 1)
        if left is not None:
            fill_bullets(left, s["bullets"])
        remove_placeholder(slide, 2)
        if not s["bullets"] and left is not None:
            remove_placeholder(slide, 1)
        set_notes(slide, s["notes"])

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    print(
        f"\nKept {existing} template slide(s), appended {len(specs)} content slide(s)."
    )
    print(f"Saved: {OUTPUT}")
    return OUTPUT


def main() -> int:
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    if not TEMPLATE.exists():
        print(f"Template missing: {TEMPLATE}", file=sys.stderr)
        return 1
    if not CONTENT.exists():
        print(f"Content missing: {CONTENT}", file=sys.stderr)
        return 1
    merge()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
