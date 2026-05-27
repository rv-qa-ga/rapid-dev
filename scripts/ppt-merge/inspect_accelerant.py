"""Inspect the Accelerant template's layouts and existing slides."""
import io
import sys
from pathlib import Path
from pptx import Presentation

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

TMPL = Path(r"C:\Automation\docs\ppt\Accelerant_PPT Training.pptx")
prs = Presentation(str(TMPL))
print(f"Template: {TMPL.name}")
print(f"Slide size: {prs.slide_width/914400:.2f} x {prs.slide_height/914400:.2f} in")
print(f"Existing slides: {len(prs.slides)}")
print(f"Slide masters: {len(prs.slide_masters)}")

for mi, master in enumerate(prs.slide_masters):
    print(f"\nMaster #{mi} ({len(master.slide_layouts)} layouts):")
    for li, layout in enumerate(master.slide_layouts):
        phs = [(ph.placeholder_format.idx, str(ph.placeholder_format.type), ph.name)
               for ph in layout.placeholders]
        print(f"  [{li:02d}] name={layout.name!r}")
        for ph in phs:
            print(f"        ph idx={ph[0]} type={ph[1]} name={ph[2]!r}")

print("\nExisting slides:")
for si, slide in enumerate(prs.slides):
    title = ""
    try:
        if slide.shapes.title and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text[:90]
    except Exception:
        pass
    print(f"  [{si:02d}] layout={slide.slide_layout.name!r:<35} title={title!r}")
