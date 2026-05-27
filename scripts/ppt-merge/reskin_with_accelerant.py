"""Re-skin the merged deck into the Accelerant training template (C-Suite polish).

Reads every slide's title / body bullets / speaker notes from the previously
merged deck, then builds a fresh deck starting from the Accelerant template
(all demo slides removed) using Accelerant's layouts with professional
typography overrides (bullets, sizing, spacing, top-anchored body).
"""
from __future__ import annotations

import io
import re
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SOURCE = Path(r"C:\Automation\reports\Scaling Quality at BSG_merged.pptx")
TEMPLATE = Path(r"C:\Automation\docs\ppt\Accelerant_PPT Training.pptx")
OUTPUT = Path(r"C:\Automation\reports\Accelerant_Scaling Quality at BSG.pptx")

TITLE_LAYOUT_NAME = "1_Title Slide"
BODY_LAYOUT_NAME = "Title and Content"
CLOSING_LAYOUT_NAME = "5_Title Slide"  # just CENTER_TITLE, for the Discussion slide

# ------- Typography -------
BULLET_CHAR_L0 = "\u2022"  # bullet
BULLET_CHAR_L1 = "\u2013"  # en dash
SIZE_L0 = Pt(20)
SIZE_L1 = Pt(16)
SIZE_SUBHEAD = Pt(22)
SPACE_AFTER_L0 = Pt(6)
SPACE_AFTER_L1 = Pt(3)
SPACE_AFTER_SUBHEAD = Pt(8)
SPACE_AFTER_SUBHEAD_FIRST = Pt(8)

# ------- Body geometry (overrides layout's right-half placeholder) -------
BODY_LEFT = Inches(0.6)
BODY_TOP = Inches(1.55)
BODY_WIDTH = Inches(14.0)
BODY_HEIGHT = Inches(6.5)

# ------- Title slide subtitle geometry tweaks -------
TITLE_SUBTITLE_SIZE = Pt(22)


def _is_title_shape(shape) -> bool:
    if not getattr(shape, "is_placeholder", False):
        return False
    t = shape.placeholder_format.type
    if t is None:
        return False
    return "TITLE" in str(t)


def extract_slides(src_path: Path) -> list[dict]:
    """Pull title / body bullets (level, text) / notes from each slide."""
    prs = Presentation(str(src_path))
    specs: list[dict] = []
    for si, slide in enumerate(prs.slides):
        title = ""
        for shape in slide.shapes:
            if _is_title_shape(shape) and shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                if txt:
                    title = txt
                    break
        if not title and slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()

        bullets: list[tuple[int, str]] = []
        seen_title = False
        for shape in slide.shapes:
            if _is_title_shape(shape):
                continue
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = para.text.replace("\n", " ").strip()
                if not text:
                    continue
                if text == title and not seen_title:
                    seen_title = True
                    continue
                if _is_noise(text):
                    continue
                bullets.append((para.level or 0, text))

        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        specs.append({"idx": si, "title": title, "bullets": bullets, "notes": notes})
    return specs


def _is_noise(text: str) -> bool:
    """Filter out layout-artifact fragments (lone digits only)."""
    t = text.strip()
    if re.fullmatch(r"\d+\.?", t):
        return True
    return False


def _is_subhead(text: str) -> bool:
    t = text.strip()
    if re.fullmatch(r"Q[1-4]", t):
        return True
    if len(t) <= 60 and t.endswith(":"):
        return True
    return False


def find_layout(prs: Presentation, name: str, fallback: str | None = None):
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            if layout.name == name:
                return layout
    if fallback:
        return find_layout(prs, fallback)
    available = [ly.name for m in prs.slide_masters for ly in m.slide_layouts]
    raise KeyError(f"Layout {name!r} not found. Available: {available}")


def delete_all_slides(prs: Presentation) -> int:
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    for sldId in ids:
        rId = sldId.get(qn("r:id"))
        try:
            prs.part.drop_rel(rId)
        except KeyError:
            pass
        sldIdLst.remove(sldId)
    return len(ids)


def get_placeholder(slide, idx: int):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None


def find_title_placeholder(slide):
    if slide.shapes.title is not None:
        return slide.shapes.title
    return get_placeholder(slide, 0)


def find_body_placeholder(slide, title_ph):
    for ph in slide.placeholders:
        if ph is title_ph:
            continue
        t = ph.placeholder_format.type
        if t is None:
            continue
        name = str(t)
        if "OBJECT" in name or "BODY" in name or "CONTENT" in name:
            return ph
    return None


def set_title(slide, text: str) -> None:
    ph = find_title_placeholder(slide)
    if ph is not None and ph.has_text_frame:
        ph.text_frame.text = text


def _set_bullet_char(paragraph, char: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum", "a:buNone", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    buFont = etree.SubElement(pPr, qn("a:buFont"))
    buFont.set("typeface", "Arial")
    buFont.set("panose", "020B0604020202020204")
    buFont.set("pitchFamily", "34")
    buFont.set("charset", "0")
    buChar = etree.SubElement(pPr, qn("a:buChar"))
    buChar.set("char", char)


def _set_no_bullet(paragraph) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum", "a:buNone", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    etree.SubElement(pPr, qn("a:buNone"))


def _set_indent(paragraph, level: int) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    if level == 0:
        pPr.set("marL", str(Inches(0.3).emu))
        pPr.set("indent", str(-Inches(0.25).emu))
    else:
        pPr.set("marL", str(Inches(0.75).emu))
        pPr.set("indent", str(-Inches(0.25).emu))


def _style_run(run, size: Pt, bold: bool = False) -> None:
    run.font.size = size
    if bold:
        run.font.bold = True


def style_body(tf, bullets: list[tuple[int, str]]) -> None:
    """Apply professional typography to the body text frame."""
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    try:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    except Exception:
        pass

    tf.clear()

    if not bullets:
        return

    for i, (stored_level, text) in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if p.runs:
            for r in p.runs:
                r.text = ""
        p.text = ""
        run = p.add_run()
        run.text = text

        is_sub = _is_subhead(text)
        effective_level = 0 if is_sub else stored_level
        p.level = effective_level

        if is_sub:
            _set_no_bullet(p)
            _set_indent(p, 0)
            p.space_before = Pt(8) if i > 0 else Pt(0)
            p.space_after = SPACE_AFTER_SUBHEAD
            _style_run(run, SIZE_SUBHEAD, bold=True)
        else:
            char = BULLET_CHAR_L0 if effective_level == 0 else BULLET_CHAR_L1
            _set_bullet_char(p, char)
            _set_indent(p, effective_level)
            p.space_after = (
                SPACE_AFTER_L0 if effective_level == 0 else SPACE_AFTER_L1
            )
            _style_run(
                run, SIZE_L0 if effective_level == 0 else SIZE_L1
            )


def reposition(shape, left: Emu, top: Emu, width: Emu, height: Emu) -> None:
    shape.left = left
    shape.top = top
    shape.width = width
    shape.height = height


def remove_unused_placeholders(slide, keep_idxs: set[int]) -> None:
    to_remove = []
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx in keep_idxs:
            continue
        t = str(ph.placeholder_format.type or "")
        if any(k in t for k in ("FOOTER", "DATE", "SLIDE_NUMBER")):
            continue
        to_remove.append(ph)
    for ph in to_remove:
        el = ph._element
        el.getparent().remove(el)


def set_notes(slide, text: str) -> None:
    if not text:
        return
    slide.notes_slide.notes_text_frame.text = text


def _classify_slide_kind(i: int, total: int, spec: dict) -> str:
    if i == 0:
        return "title"
    if i == total - 1 and not spec["bullets"]:
        return "closing"
    return "body"


def _build_title_slide(prs, spec: dict):
    layout = find_layout(prs, TITLE_LAYOUT_NAME)
    slide = prs.slides.add_slide(layout)
    title_ph = find_title_placeholder(slide)
    set_title(slide, spec["title"])

    kept = set()
    if title_ph is not None:
        kept.add(title_ph.placeholder_format.idx)

    subtitle_ph = get_placeholder(slide, 1)
    if subtitle_ph is not None and subtitle_ph.has_text_frame:
        if spec["bullets"]:
            subtitle_text = " \u2022 ".join(b[1] for b in spec["bullets"])
            subtitle_ph.text_frame.text = subtitle_text
            for p in subtitle_ph.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = TITLE_SUBTITLE_SIZE
            kept.add(1)

    set_notes(slide, spec["notes"])
    remove_unused_placeholders(slide, kept)
    return slide


def _build_body_slide(prs, spec: dict):
    layout = find_layout(prs, BODY_LAYOUT_NAME)
    slide = prs.slides.add_slide(layout)
    title_ph = find_title_placeholder(slide)
    set_title(slide, spec["title"])

    kept = set()
    if title_ph is not None:
        kept.add(title_ph.placeholder_format.idx)

    body_ph = find_body_placeholder(slide, title_ph)
    if body_ph is not None:
        reposition(body_ph, BODY_LEFT, BODY_TOP, BODY_WIDTH, BODY_HEIGHT)
        style_body(body_ph.text_frame, spec["bullets"])
        if spec["bullets"]:
            kept.add(body_ph.placeholder_format.idx)

    set_notes(slide, spec["notes"])
    remove_unused_placeholders(slide, kept)
    return slide


def _build_closing_slide(prs, spec: dict):
    try:
        layout = find_layout(prs, CLOSING_LAYOUT_NAME)
    except KeyError:
        layout = find_layout(prs, BODY_LAYOUT_NAME)
    slide = prs.slides.add_slide(layout)
    set_title(slide, spec["title"])
    kept = set()
    title_ph = find_title_placeholder(slide)
    if title_ph is not None:
        kept.add(title_ph.placeholder_format.idx)
        if title_ph.has_text_frame:
            tf = title_ph.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            for p in tf.paragraphs:
                p.alignment = PP_ALIGN.CENTER
    remove_unused_placeholders(slide, kept)
    return slide


def build(specs: list[dict]) -> Path:
    prs = Presentation(str(TEMPLATE))
    n_removed = delete_all_slides(prs)
    print(f"Removed {n_removed} demo slide(s) from Accelerant template")

    total = len(specs)
    for i, s in enumerate(specs):
        kind = _classify_slide_kind(i, total, s)
        if kind == "title":
            _build_title_slide(prs, s)
        elif kind == "closing":
            _build_closing_slide(prs, s)
        else:
            _build_body_slide(prs, s)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUTPUT))
    print(f"Built {total} slide(s) with C-suite styling")
    print(f"Saved: {OUTPUT}")
    return OUTPUT


def main() -> int:
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    if not SOURCE.exists():
        print(f"Source missing: {SOURCE}", file=sys.stderr)
        return 1
    if not TEMPLATE.exists():
        print(f"Template missing: {TEMPLATE}", file=sys.stderr)
        return 1

    specs = extract_slides(SOURCE)
    print(f"Extracted {len(specs)} slide(s) from source:")
    for s in specs:
        sub_n = sum(1 for lvl, _ in s["bullets"] if lvl > 0)
        sub_h = sum(1 for _, t in s["bullets"] if _is_subhead(t))
        print(
            f"  [{s['idx']:02d}] {s['title']!r:<55} "
            f"total={len(s['bullets']):>2} sub_heads={sub_h} lvl1={sub_n} "
            f"notes={'Y' if s['notes'] else '-'}"
        )

    build(specs)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
