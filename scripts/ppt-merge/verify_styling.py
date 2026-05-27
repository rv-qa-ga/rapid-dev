"""Verify the bullets/styling applied to slides in the polished Accelerant output."""
import io
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Emu
from lxml import etree
from pptx.oxml.ns import qn

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

OUT = Path(r"C:\Automation\reports\Accelerant_Scaling Quality at BSG.pptx")
prs = Presentation(str(OUT))
print(f"Slides: {len(prs.slides)}")

for si in [3, 8, 13, 15]:  # Agenda, roadmap, loop diagram, Discussion
    slide = prs.slides[si]
    print("\n" + "=" * 60)
    print(f"Slide [{si:02d}] layout={slide.slide_layout.name!r}")
    for sh in slide.shapes:
        if not sh.is_placeholder or not sh.has_text_frame:
            continue
        pt = sh.placeholder_format.type
        if pt is None or "TITLE" in str(pt):
            continue
        tf = sh.text_frame
        print(f"  placeholder idx={sh.placeholder_format.idx} type={pt}")
        print(f"    geom: L={Emu(sh.left).inches:.2f} T={Emu(sh.top).inches:.2f} "
              f"W={Emu(sh.width).inches:.2f} H={Emu(sh.height).inches:.2f} in")
        print(f"    anchor={tf.vertical_anchor} auto_size={tf.auto_size}")
        for pi, p in enumerate(tf.paragraphs):
            txt = p.text[:70]
            size = None
            bold = None
            for r in p.runs:
                if r.font.size:
                    size = r.font.size.pt
                if r.font.bold is not None:
                    bold = r.font.bold
                break
            pPr = p._p.find(qn("a:pPr"))
            bu = "inherit"
            if pPr is not None:
                if pPr.find(qn("a:buNone")) is not None:
                    bu = "none"
                elif pPr.find(qn("a:buChar")) is not None:
                    bu = f"char:{pPr.find(qn('a:buChar')).get('char')}"
            print(f"    p[{pi}] lvl={p.level} sz={size} b={bold} bullet={bu} text={txt!r}")
