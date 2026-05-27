import io
import sys
from pathlib import Path
from pptx import Presentation

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

OUT = Path(r"C:\Automation\reports\Accelerant_Scaling Quality at BSG.pptx")
prs = Presentation(str(OUT))
for i, slide in enumerate(prs.slides):
    print(f"\nSlide [{i:02d}] layout={slide.slide_layout.name!r}")
    print(f"  slide.shapes.title = {slide.shapes.title!r}")
    for j, sh in enumerate(slide.shapes):
        kind = type(sh).__name__
        is_ph = sh.is_placeholder
        ph_info = ""
        if is_ph:
            ph_info = f" [ph idx={sh.placeholder_format.idx} type={sh.placeholder_format.type}]"
        has_tf = sh.has_text_frame
        preview = ""
        if has_tf:
            preview = sh.text_frame.text.replace("\n", " | ")[:100]
        print(f"    shape[{j}] {kind}{ph_info} name={sh.name!r} text={preview!r}")
    if i == 15:
        break
