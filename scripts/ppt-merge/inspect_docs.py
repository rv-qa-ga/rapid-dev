"""Inspect template PPT and Word content doc to understand structure before merging."""
from __future__ import annotations

import sys
import io
from pathlib import Path
from collections import Counter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from pptx import Presentation
from docx import Document

TEMPLATE = Path(r"C:\Automation\docs\ppt\Scaling Quality at BSG.pptx")
CONTENT = Path(r"C:\Automation\docs\ppt\Offsite Presentation Content.docx")


def inspect_template(path: Path) -> None:
    print("=" * 70)
    print(f"TEMPLATE: {path.name}")
    print("=" * 70)
    prs = Presentation(str(path))
    print(f"Slide size: {prs.slide_width} x {prs.slide_height} EMU "
          f"({prs.slide_width/914400:.2f} x {prs.slide_height/914400:.2f} in)")
    print(f"Total slides in template: {len(prs.slides)}")
    print(f"Slide masters: {len(prs.slide_masters)}")
    for mi, master in enumerate(prs.slide_masters):
        print(f"  Master #{mi}: {len(master.slide_layouts)} layouts")
        for li, layout in enumerate(master.slide_layouts):
            phs = [(ph.placeholder_format.idx, ph.placeholder_format.type, ph.name)
                   for ph in layout.placeholders]
            print(f"    Layout[{li}] name={layout.name!r} placeholders={phs}")
    print("\nExisting slides in template:")
    for si, slide in enumerate(prs.slides):
        title = ""
        try:
            if slide.shapes.title and slide.shapes.title.has_text_frame:
                title = slide.shapes.title.text_frame.text[:80]
        except Exception:
            pass
        print(f"  Slide[{si}] layout={slide.slide_layout.name!r} title={title!r}")


def inspect_docx(path: Path) -> None:
    print("\n" + "=" * 70)
    print(f"CONTENT: {path.name}")
    print("=" * 70)
    doc = Document(str(path))
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    print(f"Total tables: {len(doc.tables)}")
    style_counts = Counter(p.style.name for p in doc.paragraphs if p.text.strip())
    print(f"Paragraph styles (non-empty): {dict(style_counts)}")

    print("\nAll non-empty paragraphs (idx | style | text):")
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if not txt:
            continue
        print(f"  {i:03d} [{p.style.name}] {txt}")

    # count images
    img_count = 0
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            img_count += 1
    print(f"\nEmbedded images (approx): {img_count}")


def main() -> int:
    if not TEMPLATE.exists():
        print(f"Template missing: {TEMPLATE}", file=sys.stderr)
        return 1
    if not CONTENT.exists():
        print(f"Content missing: {CONTENT}", file=sys.stderr)
        return 1
    inspect_template(TEMPLATE)
    inspect_docx(CONTENT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
